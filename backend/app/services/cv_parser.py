from pathlib import Path

import pymupdf
from docx import Document

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_DOCUMENT_PAGES = 80
MAX_EXTRACTED_CHARACTERS = 60_000


class CVParserError(Exception):
    """Raised when a CV cannot be parsed safely."""


def _check_length(text: str) -> str:
    text = text.strip()
    if len(text) > MAX_EXTRACTED_CHARACTERS:
        raise CVParserError(
            "The extracted document is too long for this version of CareerMate "
            f"({len(text):,} characters). Please use a shorter CV (maximum "
            f"{MAX_EXTRACTED_CHARACTERS:,} characters)."
        )
    return text


def extract_pdf_text(file_path: Path) -> str:
    try:
        with pymupdf.open(file_path) as document:
            if document.is_encrypted and not document.authenticate(""):
                raise CVParserError(
                    "This PDF is password-protected. Remove its password and try again."
                )
            if len(document) > MAX_DOCUMENT_PAGES:
                raise CVParserError(
                    f"The PDF has too many pages. Maximum: {MAX_DOCUMENT_PAGES}."
                )
            text = "\n\n".join(page.get_text("text") for page in document)
            return _check_length(text)
    except CVParserError:
        raise
    except Exception as exc:
        raise CVParserError(
            "Could not read this PDF. Check that it is a valid, uncorrupted PDF."
        ) from exc


def extract_docx_text(file_path: Path) -> str:
    try:
        document = Document(file_path)
        sections: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                sections.append(text)

        # CVs often put skills or employment history in tables.
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                row_text = " | ".join(cell for cell in cells if cell)
                if row_text:
                    sections.append(row_text)

        return _check_length("\n".join(sections))
    except CVParserError:
        raise
    except Exception as exc:
        raise CVParserError(
            "Could not read this DOCX. Check that it is a valid Word document."
        ) from exc


def extract_txt_text(file_path: Path) -> str:
    try:
        return _check_length(
            file_path.read_text(encoding="utf-8-sig", errors="replace")
        )
    except CVParserError:
        raise
    except Exception as exc:
        raise CVParserError("Could not read this text file.") from exc


def extract_cv_text(file_path: Path) -> str:
    extension = file_path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise CVParserError("Supported CV formats are PDF, DOCX, and TXT.")

    if extension == ".pdf":
        return extract_pdf_text(file_path)
    if extension == ".docx":
        return extract_docx_text(file_path)
    if extension == ".txt":
        return extract_txt_text(file_path)

    raise CVParserError("Unsupported CV format.")
